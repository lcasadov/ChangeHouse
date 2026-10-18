#!/usr/bin/env python3
"""Genera las plantillas Word de los entregables de ChangeHouse.

Lee los bloques ``ficha``, ``aprobacion`` e ``indice`` de cada prompt
``.claude/commands/harness/entregable-*.md`` y escribe la plantilla
correspondiente en ``.claude/harness/profile/entregables/``, con la identidad
visual descrita en ``formato.md`` de esa misma carpeta. También dibuja el
logotipo y la banda de portada en ``assets/``.

Uso:
    python .claude/harness/scripts/build-entregables.py
    python .claude/harness/scripts/build-entregables.py --only funcional

Requiere python-docx y Pillow.
"""
from __future__ import annotations

import argparse
import math
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[3]
CMD_DIR = ROOT / ".claude" / "commands" / "harness"
OUT_DIR = ROOT / ".claude" / "harness" / "profile" / "entregables"
ASSETS = OUT_DIR / "assets"

# --- Identidad (ver formato.md) -------------------------------------------
CORAL = "E8505B"
NAVY = "1F2A44"
TINT = "FDEEEF"
TEXT = "3A3A3A"
MUTED = "5A6672"
BORDER = "D9DDE3"
WHITE = "FFFFFF"
FONT = "Calibri"
MONO = "Consolas"
TABLE_WIDTH_CM = 16.0


@dataclass(frozen=True)
class Entregable:
    clave: str
    prompt: str
    salida: str
    tipo: str
    primera_emision: bool = False


ENTREGABLES = [
    Entregable("funcional", "entregable-documento-funcional.md",
               "documento-funcional.docx", "Documento funcional"),
    Entregable("arquitectura", "entregable-documento-arquitectura.md",
               "documento-arquitectura.docx", "Documento de arquitectura"),
    Entregable("seguridad-primera", "entregable-informe-seguridad-primera-emision.md",
               "informe-seguridad-primera-emision.docx",
               "Informe de seguridad SAST y SCA", primera_emision=True),
    Entregable("seguridad-reemision", "entregable-informe-seguridad-reemision.md",
               "informe-seguridad-reemision.docx",
               "Informe de seguridad SAST y SCA — Reemisión"),
]

FICHA_MARCADORES = {
    "Sistema": "{{SISTEMA}}",
    "Organización": "{{ORGANIZACION}}",
    "Estado": "{{ESTADO}}",
    "Clasificación": "{{CLASIFICACION}}",
    "Revisión": "{{REVISION}}",
    "Fecha": "{{FECHA}}",
    "Fecha del informe": "{{FECHA}}",
    "Tipo de prueba": "Análisis estático de código (SAST) y análisis de composición de software (SCA)",
}


# --- Lectura de los prompts -------------------------------------------------
def bloque(texto: str, nombre: str, prompt: str) -> list[str]:
    m = re.search(rf"^```{nombre}\n(.*?)^```", texto, re.S | re.M)
    if not m:
        sys.exit(f"Falta el bloque ```{nombre} en {prompt}")
    return [l.rstrip() for l in m.group(1).splitlines() if l.strip()]


def nivel(linea: str) -> tuple[int, str]:
    t = linea.strip()
    if re.match(r"^Anexo [A-Z] — ", t) or re.match(r"^\d+\.\s", t):
        return 1, t
    if re.match(r"^(\d+|[A-Z])\.(\d+|n)\s", t):
        return 2, t
    if re.match(r"^\d+\.(\d+|n)\.\d+\s", t):
        return 3, t
    sys.exit(f"Línea de índice no reconocida: {linea!r}")


# --- Imágenes ---------------------------------------------------------------
def _fuente(tamano: int, negrita: bool) -> ImageFont.FreeTypeFont:
    candidatos = (["segoeuib.ttf", "calibrib.ttf", "arialbd.ttf", "DejaVuSans-Bold.ttf"]
                  if negrita else ["segoeui.ttf", "calibri.ttf", "arial.ttf", "DejaVuSans.ttf"])
    for nombre in candidatos:
        for base in (Path("C:/Windows/Fonts"), Path("/Library/Fonts"), Path("/usr/share/fonts/truetype/dejavu")):
            ruta = base / nombre
            if ruta.exists():
                return ImageFont.truetype(str(ruta), tamano)
    return ImageFont.load_default(tamano)


def _rgb(hexa: str) -> tuple[int, int, int]:
    return tuple(int(hexa[i:i + 2], 16) for i in (0, 2, 4))


def _arco_con_flecha(d: ImageDraw.ImageDraw, centro: tuple[int, int], radio: int, desde: float,
                     hasta: float, grosor: int, color) -> None:
    """Arco en sentido horario de ``desde`` a ``hasta`` (grados, 0 = las 3 en punto) con punta final."""
    cx, cy = centro
    d.arc([cx - radio, cy - radio, cx + radio, cy + radio], desde, hasta, fill=color, width=grosor)
    theta = math.radians(hasta)
    medio = radio - grosor / 2  # el arco se dibuja hacia dentro del rectángulo
    px, py = cx + medio * math.cos(theta), cy + medio * math.sin(theta)
    radial = (math.cos(theta), math.sin(theta))
    tangente = (-math.sin(theta), math.cos(theta))
    largo, ancho = grosor * 1.6, grosor * 1.25
    d.polygon([
        (px + tangente[0] * largo, py + tangente[1] * largo),
        (px + radial[0] * ancho, py + radial[1] * ancho),
        (px - radial[0] * ancho, py - radial[1] * ancho),
    ], fill=color)


def dibujar_logo(ruta: Path) -> None:
    """Casa blanca dentro de un ciclo de dos flechas: cambio de casa y cadenas de intercambio."""
    s = 2048  # se dibuja a 4x y se reduce para suavizar los bordes
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle([0, 0, s - 1, s - 1], radius=440, fill=_rgb(CORAL))

    navy, blanco = _rgb(NAVY), _rgb(WHITE)
    centro, radio, grosor = (1024, 1040), 760, 125
    _arco_con_flecha(d, centro, radio, 205, 320, grosor, navy)
    _arco_con_flecha(d, centro, radio, 25, 140, grosor, navy)

    d.polygon([(1024, 520), (600, 900), (1448, 900)], fill=blanco)          # tejado
    d.rectangle([720, 900, 1328, 1420], fill=blanco)                         # cuerpo
    d.rectangle([950, 1150, 1098, 1420], fill=_rgb(CORAL))                   # puerta
    d.rectangle([780, 990, 900, 1100], fill=_rgb(CORAL))                     # ventana
    d.rectangle([1148, 990, 1268, 1100], fill=_rgb(CORAL))                   # ventana
    img.resize((512, 512), Image.LANCZOS).save(ruta)


def dibujar_portada(ruta: Path, logo: Path) -> None:
    ancho, alto = 2400, 900
    img = Image.new("RGB", (ancho, alto))
    a, b = _rgb(NAVY), _rgb(CORAL)
    d = ImageDraw.Draw(img)
    for x in range(ancho):
        t = x / (ancho - 1)
        d.line([(x, 0), (x, alto)], fill=tuple(int(a[i] + (b[i] - a[i]) * t) for i in range(3)))
    marca = Image.open(logo).resize((480, 480))
    img.paste(marca, (140, 210), marca)
    d.text((720, 250), "ChangeHouse", font=_fuente(200, True), fill=_rgb(WHITE))
    d.text((730, 520), "Cambia de casa sin vender primero", font=_fuente(72, False), fill=_rgb(WHITE))
    img.save(ruta)


# --- Utilidades de Word -----------------------------------------------------
def _fijar_fuente(rpr_owner, nombre: str) -> None:
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for atributo in list(fonts.attrib):
        if atributo.endswith("Theme"):
            del fonts.attrib[atributo]
    for clave in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(clave), nombre)


def _quitar_color_tema(style) -> None:
    color = style.element.rPr.find(qn("w:color")) if style.element.rPr is not None else None
    if color is not None:
        for atributo in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(atributo), None)


def estilo_texto(style, tamano: float, color: str, negrita=False, cursiva=False) -> None:
    style.font.name = FONT
    _fijar_fuente(style.element, FONT)
    style.font.size = Pt(tamano)
    style.font.bold = negrita
    style.font.italic = cursiva
    style.font.color.rgb = RGBColor.from_string(color)
    _quitar_color_tema(style)


def borde_inferior(parrafo_o_estilo, color: str, tamano: int = 12) -> None:
    ppr = parrafo_o_estilo.element.get_or_add_pPr() if hasattr(parrafo_o_estilo.element, "get_or_add_pPr") \
        else parrafo_o_estilo._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(tamano)), ("w:space", "4"), ("w:color", color)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def configurar_estilos(doc: Document) -> None:
    st = doc.styles
    normal = st["Normal"]
    estilo_texto(normal, 10.5, TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nombre, tamano, color, antes in (("Heading 1", 18, NAVY, 0), ("Heading 2", 14, CORAL, 14),
                                         ("Heading 3", 12, NAVY, 10)):
        h = st[nombre]
        estilo_texto(h, tamano, color, negrita=True)
        h.paragraph_format.space_before = Pt(antes)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
    st["Heading 1"].paragraph_format.page_break_before = True
    st["Heading 1"].paragraph_format.space_after = Pt(12)
    borde_inferior(st["Heading 1"], CORAL)

    cap = st["Caption"]
    estilo_texto(cap, 9, MUTED, cursiva=True)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for nombre in ("List Bullet", "List Number"):
        estilo_texto(st[nombre], 10.5, TEXT)

    pre = st.add_style("Título preliminar", WD_STYLE_TYPE.PARAGRAPH)
    pre.base_style = normal
    estilo_texto(pre, 16, NAVY, negrita=True)
    pre.paragraph_format.space_before = Pt(6)
    pre.paragraph_format.space_after = Pt(10)
    pre.paragraph_format.keep_with_next = True
    borde_inferior(pre, CORAL)

    ind = st.add_style("Indicación", WD_STYLE_TYPE.PARAGRAPH)
    ind.base_style = normal
    estilo_texto(ind, 10.5, MUTED, cursiva=True)


def configurar_pagina(doc: Document) -> None:
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.5)
    s.top_margin, s.bottom_margin = Cm(2.5), Cm(2.2)
    s.footer_distance = Cm(1.0)
    s.different_first_page_header_footer = True  # la portada no lleva pie


def _sombrear(celda, color: str) -> None:
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def _ancho_celda(celda, cm: float) -> None:
    celda.width = Cm(cm)
    tcpr = celda._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(cm * 567)))
    tcw.set(qn("w:type"), "dxa")


def _diseno_fijo(tabla) -> None:
    tblpr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    tabla.autofit = False


def _bordes_tabla(tabla, color: str, ninguno=False) -> None:
    tblpr = tabla._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        if ninguno:
            e.set(qn("w:val"), "nil")
        else:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), color)
        borders.append(e)
    tblpr.append(borders)


def _texto_celda(celda, texto: str, negrita=False, color=TEXT, tamano=9.0,
                 alineacion=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = alineacion
    run = p.add_run(texto)
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    run.font.color.rgb = RGBColor.from_string(color)
    if texto.startswith("[["):
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string(MUTED)


def tabla_datos(doc: Document, cabecera: list[str], filas: list[list[str]], anchos: list[float]) -> None:
    t = doc.add_table(rows=1 + len(filas), cols=len(cabecera))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _diseno_fijo(t)
    _bordes_tabla(t, BORDER)
    for i, titulo in enumerate(cabecera):
        c = t.rows[0].cells[i]
        _ancho_celda(c, anchos[i])
        _sombrear(c, NAVY)
        _texto_celda(c, titulo, negrita=True, color=WHITE)
    for r, fila in enumerate(filas, start=1):
        for i, valor in enumerate(fila):
            c = t.rows[r].cells[i]
            _ancho_celda(c, anchos[i])
            if r % 2 == 0:
                _sombrear(c, TINT)
            _texto_celda(c, valor, negrita=(i == 0 and len(cabecera) == 2))
    # Repite la cabecera si la tabla salta de página.
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    doc.add_paragraph()


def _campo(parrafo, instruccion: str, marcador: str = "", tamano: float | None = None,
           color: str | None = None) -> None:
    def run_con(elemento):
        r = parrafo.add_run()
        if tamano:
            r.font.size = Pt(tamano)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        r._r.append(elemento)
        return r
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    inicio.set(qn("w:dirty"), "true")
    run_con(inicio)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    run_con(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run_con(sep)
    r = parrafo.add_run(marcador)
    if tamano:
        r.font.size = Pt(tamano)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run_con(fin)


def _actualizar_campos_al_abrir(doc: Document) -> None:
    settings = doc.settings.element
    e = OxmlElement("w:updateFields")
    e.set(qn("w:val"), "true")
    settings.append(e)


# --- Partes del documento ---------------------------------------------------
def portada(doc: Document, ent: Entregable, banda: Path) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(banda), width=Cm(16))
    for _ in range(3):
        doc.add_paragraph()

    def linea(texto, tamano, color, negrita=True, despues=8):
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(despues)
        r = q.add_run(texto)
        r.font.size, r.font.bold = Pt(tamano), negrita
        r.font.color.rgb = RGBColor.from_string(color)

    linea("{{SISTEMA}}", 30, NAVY)
    linea("{{DESCRIPCION}}", 16, MUTED, negrita=False, despues=28)
    linea(ent.tipo, 15, CORAL)
    for _ in range(6):
        doc.add_paragraph()
    linea("Revisión {{REVISION}} · {{FECHA}}", 11, MUTED, negrita=False, despues=2)
    linea("{{ORGANIZACION}} · {{CLASIFICACION}}", 11, MUTED, negrita=False)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def preliminares(doc: Document, ent: Entregable, ficha: list[str], aprobacion: list[str]) -> None:
    doc.add_paragraph("Ficha del documento", style="Título preliminar")
    filas = []
    for campo in ficha:
        if campo == "Título":
            valor = f"{ent.tipo} — {{{{SISTEMA}}}}"
        else:
            valor = FICHA_MARCADORES.get(campo, f"[[{campo}]]")
        filas.append([campo, valor])
    tabla_datos(doc, ["Campo", "Valor"], filas, [5.0, 11.0])

    doc.add_paragraph("Histórico de versiones", style="Título preliminar")
    descripcion = "Primera emisión" if ent.primera_emision else "[[Descripción de los cambios de esta revisión]]"
    tabla_datos(doc, ["Revisión", "Fecha", "Autor", "Descripción"],
                [["{{REVISION}}", "{{FECHA}}", "{{AUTOR}}", descripcion]], [2.2, 2.8, 4.0, 7.0])

    titulo = doc.add_paragraph("Revisión y aprobación", style="Título preliminar")
    titulo.paragraph_format.page_break_before = True
    filas = [[rol, "{{AUTOR}}" if rol == "Autor" else "[[Nombre]]",
              "{{FECHA}}" if rol == "Autor" else "[[Fecha]]", ""] for rol in aprobacion]
    tabla_datos(doc, ["Rol", "Nombre", "Fecha", "Firma"], filas, [4.5, 5.0, 2.8, 3.7])

    indice = doc.add_paragraph("Índice", style="Título preliminar")
    indice.paragraph_format.page_break_before = True
    _campo(doc.add_paragraph(), r'TOC \o "1-3" \h \z \u',
           "Actualiza el índice: clic derecho sobre él y «Actualizar campo».")


def cuerpo(doc: Document, indice: list[str]) -> None:
    entradas = [nivel(l) for l in indice]
    for i, (n, texto) in enumerate(entradas):
        repetible = re.match(r"^(.*?) — \((.*)\)$", texto)
        if repetible:
            titulo = f"{repetible.group(1)} — [[nombre]]"
            indicacion = (f"[[Apartado repetible: {repetible.group(2)}. Duplica título e indicación "
                          f"tantas veces como haga falta y renumera los apartados siguientes.]]")
        else:
            titulo = texto
            nombre = re.sub(r"^(Anexo [A-Z] — |(\d+|[A-Z])(\.(\d+|n))*\.?\s)", "", texto)
            indicacion = f"[[Redacta «{nombre}» según las indicaciones del prompt.]]"
        doc.add_paragraph(titulo, style=f"Heading {n}")
        tiene_hijos = i + 1 < len(entradas) and entradas[i + 1][0] > n
        if not tiene_hijos:
            doc.add_paragraph(indicacion, style="Indicación")


def pie(doc: Document, ent: Entregable, logo: Path) -> None:
    footer = doc.sections[0].footer
    footer.is_linked_to_previous = False
    base = footer.paragraphs[0]
    anchos = [1.2, 6.4, 3.0, 2.4, 3.0]
    t = footer.add_table(rows=1, cols=5, width=Cm(TABLE_WIDTH_CM))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _diseno_fijo(t)
    _bordes_tabla(t, BORDER, ninguno=True)
    celdas = t.rows[0].cells
    for c, w in zip(celdas, anchos):
        _ancho_celda(c, w)
        tcpr = c._tc.get_or_add_tcPr()
        b = OxmlElement("w:tcBorders")
        top = OxmlElement("w:top")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:color", CORAL)):
            top.set(qn(k), v)
        b.append(top)
        tcpr.append(b)
    celdas[0].paragraphs[0].add_run().add_picture(str(logo), width=Cm(0.8))
    _texto_celda(celdas[1], f"{ent.tipo} · {{{{SISTEMA}}}}", tamano=8, color=MUTED)
    _texto_celda(celdas[2], "Rev. {{REVISION}}", tamano=8, color=MUTED, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    _texto_celda(celdas[3], "{{FECHA}}", tamano=8, color=MUTED, alineacion=WD_ALIGN_PARAGRAPH.CENTER)
    pagina = celdas[4].paragraphs[0]
    pagina.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pagina.paragraph_format.space_after = Pt(0)
    for texto, instruccion in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
        r = pagina.add_run(texto)
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        _campo(pagina, instruccion, "1", tamano=8, color=MUTED)
    # El párrafo vacío que exige Word tras la tabla, reducido a 1 pt.
    footer._element.remove(base._p)
    cierre = footer.add_paragraph()
    cierre.paragraph_format.space_after = Pt(0)
    cierre.add_run().font.size = Pt(1)


def construir(ent: Entregable, logo: Path, banda: Path) -> Path:
    texto = (CMD_DIR / ent.prompt).read_text(encoding="utf-8")
    ficha = bloque(texto, "ficha", ent.prompt)
    aprobacion = bloque(texto, "aprobacion", ent.prompt)
    indice = bloque(texto, "indice", ent.prompt)

    doc = Document()
    configurar_pagina(doc)
    configurar_estilos(doc)
    doc.core_properties.title = ent.tipo
    doc.core_properties.author = "ChangeHouse"
    doc.core_properties.comments = f"Plantilla generada desde {ent.prompt}"
    portada(doc, ent, banda)
    preliminares(doc, ent, ficha, aprobacion)
    cuerpo(doc, indice)
    pie(doc, ent, logo)
    _actualizar_campos_al_abrir(doc)
    # Quita el párrafo vacío inicial que crea python-docx.
    primero = doc.paragraphs[0]
    if not primero.text and not primero.runs:
        primero._p.getparent().remove(primero._p)

    salida = OUT_DIR / ent.salida
    doc.save(salida)
    return salida


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--only", choices=[e.clave for e in ENTREGABLES],
                        help="Genera solo la plantilla indicada")
    args = parser.parse_args()

    ASSETS.mkdir(parents=True, exist_ok=True)
    logo, banda = ASSETS / "logo.png", ASSETS / "portada.png"
    dibujar_logo(logo)
    dibujar_portada(banda, logo)

    for ent in ENTREGABLES:
        if args.only and ent.clave != args.only:
            continue
        ruta = construir(ent, logo, banda)
        print(f"ok  {ruta.relative_to(ROOT).as_posix()}")


if __name__ == "__main__":
    main()
